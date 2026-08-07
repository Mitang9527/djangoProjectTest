"""
Excel (.xlsx) 导出器
基于 openpyxl，自动列宽、表头样式、数据类型识别
"""
from io import BytesIO

from django.db.models import QuerySet
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .base import BaseExporter


class ExcelExporter(BaseExporter):
    """Excel 格式导出器"""
    content_type = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    extension = 'xlsx'

    # 样式常量
    HEADER_FONT = Font(name='微软雅黑', bold=True, size=11, color='FFFFFF')
    HEADER_FILL = PatternFill(start_color='1F2937', end_color='1F2937', fill_type='solid')
    HEADER_ALIGNMENT = Alignment(horizontal='center', vertical='center')
    CELL_FONT = Font(name='微软雅黑', size=10)
    CELL_ALIGNMENT = Alignment(vertical='center')
    THIN_BORDER = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB'),
    )

    def export(self, queryset: QuerySet) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = self.config.filename[:31]  # Excel sheet name max 31 chars

        headers = self.config.get_headers()

        # ---- 写表头 ----
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = self.HEADER_FONT
            cell.fill = self.HEADER_FILL
            cell.alignment = self.HEADER_ALIGNMENT
            cell.border = self.THIN_BORDER

        # ---- 写数据行 ----
        for row_idx, obj in enumerate(queryset.iterator(chunk_size=1000), 2):
            for col_idx, field in enumerate(self.config.fields, 1):
                value = self._resolve_value(obj, field)
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.font = self.CELL_FONT
                cell.alignment = self.CELL_ALIGNMENT
                cell.border = self.THIN_BORDER

        # ---- 自动列宽 ----
        for col_idx in range(1, len(headers) + 1):
            max_width = len(str(headers[col_idx - 1])) * 2.5  # 中文字符宽
            # 取前 100 行采样（避免遍历全部数据）
            for row in ws.iter_rows(min_col=col_idx, max_col=col_idx,
                                    min_row=2, max_row=min(101, ws.max_row),
                                    values_only=True):
                for cell_val in row:
                    if cell_val:
                        max_width = max(max_width, len(str(cell_val)) * 2.2)
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_width + 4, 60)

        # ---- 冻结首行 ----
        ws.freeze_panes = 'A2'

        # ---- 添加筛选器（如果是小结果集） ----
        if ws.max_row <= 10000:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"

        # 写入 BytesIO
        output = BytesIO()
        wb.save(output)
        return output.getvalue()
